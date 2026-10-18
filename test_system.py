"""Интеграционные проверки API, прав, атомарности и расчётов."""
from datetime import date
from decimal import Decimal
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest
import os
from unittest.mock import patch
from urllib.request import Request, urlopen
from urllib.error import HTTPError

from api import Api, ApiError
from domain import DomainError, interval, units, overlaps, password_matches, password_hash
from storage import Store
from server import make_server
from restore import restore
from maintenance import InstanceLock, verify_backup
from economics import calculate, sensitivity
import exports
import quality


class SystemTest(unittest.TestCase):
    def test_loopback_connection_with_system_proxy(self):
        proxy = 'http://127.0.0.1:1'
        with patch.dict(os.environ, {'HTTP_PROXY': proxy, 'HTTPS_PROXY': proxy,
                                    'http_proxy': proxy, 'https_proxy': proxy,
                                    'NO_PROXY': '', 'no_proxy': ''}):
            api = Api(self.url)
            user = api.login('teacher', 'DemoTime2026!')
            self.assertEqual(user['role'], 'teacher')
            api.logout()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name) / 'test.sqlite')
        self.store.demo()
        self.server = make_server(self.store, port=0)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = 'http://127.0.0.1:' + str(self.server.server_port)
        self.admin = self.client('admin')
        self.head = self.client('head')
        self.teacher = self.client('teacher')
        self.other = self.client('teacher2')
        self.day = date.today().replace(day=1).isoformat()
        self.period = self.day[:7]

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)
        self.temp.cleanup()

    def client(self, login):
        api = Api(self.url)
        api.login(login, '12345' if login == 'admin' else 'DemoTime2026!')
        return api

    def record(self, start='09:00', end='10:30', teacher_id=1):
        return {'teacher_id': teacher_id, 'workday': self.day, 'start': start, 'end': end, 'type_id': 1, 'title': 'Контрольное занятие', 'evidence': 'Тестовый пример'}

    def create(self, start='09:00', end='10:30'):
        result = self.teacher.call('entry.save', self.record(start, end))
        return self.row(result['id'])

    def row(self, ident):
        return next(r for r in self.admin.call('entries', {'period': self.period}) if r['id'] == ident)

    def send(self, api, row, target, reason=''):
        return api.call('entry.transition', {'id': row['id'], 'version': row['version'], 'target': target, 'reason': reason})

    def test_login_and_password_hash(self):
        digest = password_hash('ValidPassword23')
        self.assertNotIn('ValidPassword23', digest)
        self.assertTrue(password_matches('ValidPassword23', digest))
        self.assertFalse(password_matches('BadPassword23', digest))
        with self.assertRaises(ApiError):
            Api(self.url).login('teacher', 'incorrect')

    def test_teacher_scope_and_department_isolation(self):
        row = self.create()
        self.assertEqual(len(self.teacher.call('entries')), 1)
        self.assertEqual(self.other.call('entries'), [])
        self.assertEqual(len(self.head.call('entries')), 1)
        with self.assertRaises(ApiError) as error:
            self.teacher.call('entry.save', self.record('12:00', '13:00', teacher_id=7))
        self.assertEqual(error.exception.status, 403)
        with self.assertRaises(ApiError):
            self.other.call('entry.history', {'id': row['id']})

    def test_role_restrictions(self):
        for method in ('users', 'audit', 'backup'):
            with self.assertRaises(ApiError) as error:
                self.teacher.call(method)
            self.assertEqual(error.exception.status, 403)
        with self.assertRaises(ApiError):
            self.teacher.call('settings.save', {'organization': 'Подмена'})

    def test_intervals_and_overlap_boundary(self):
        self.assertEqual(interval('09:00', '10:30')[2], 90)
        self.assertEqual(units(90, 'academic'), '2.00')
        self.assertEqual(units(90, 'clock'), '1.50')
        self.assertFalse(overlaps(540, 600, 600, 660))
        self.create('09:00', '10:30')
        self.create('10:30', '11:00')
        with self.assertRaises(ApiError) as error:
            self.create('10:00', '11:30')
        self.assertEqual(error.exception.status, 409)

    def test_invalid_time_future_and_missing_type(self):
        for first, last in [('10:00', '09:00'), ('25:00', '26:00'), ('9:00', '10:00')]:
            with self.assertRaises(ApiError):
                self.create(first, last)
        record = self.record()
        record['workday'] = '2099-01-01'
        with self.assertRaises(ApiError):
            self.teacher.call('entry.save', record)
        record = self.record()
        record['type_id'] = 999
        with self.assertRaises(ApiError):
            self.teacher.call('entry.save', record)

    def test_agreement_and_immutable_approved_record(self):
        row = self.create()
        self.send(self.teacher, row, 'submitted')
        row = self.row(row['id'])
        with self.assertRaises(ApiError):
            self.send(self.teacher, row, 'approved')
        self.send(self.head, row, 'approved')
        row = self.row(row['id'])
        with self.assertRaises(ApiError):
            self.teacher.call('entry.save', row)
        self.assertEqual(len(self.teacher.call('entry.history', {'id': row['id']})), 2)

    def test_return_requires_reason_and_resubmit(self):
        row = self.create()
        self.send(self.teacher, row, 'submitted')
        row = self.row(row['id'])
        with self.assertRaises(ApiError):
            self.send(self.head, row, 'returned')
        self.send(self.head, row, 'returned', 'Уточнить основание')
        row = self.row(row['id'])
        row['evidence'] = 'Исправленное основание'
        self.teacher.call('entry.save', row)
        self.send(self.teacher, self.row(row['id']), 'submitted')
        self.assertEqual(self.row(row['id'])['status'], 'submitted')

    def test_optimistic_lock(self):
        row = self.create()
        copy = dict(row)
        row['title'] = 'Новая версия'
        self.teacher.call('entry.save', row)
        copy['title'] = 'Устаревшая версия'
        with self.assertRaises(ApiError) as error:
            self.teacher.call('entry.save', copy)
        self.assertEqual(error.exception.status, 409)
        self.assertEqual(self.row(row['id'])['title'], 'Новая версия')

    def test_report_only_approved_and_separate_units(self):
        row = self.create()
        self.send(self.teacher, row, 'submitted')
        report = self.teacher.call('report', {'period': self.period})
        self.assertEqual(report['approved_minutes'], 0)
        self.assertEqual(report['submitted_minutes'], 90)
        self.send(self.head, self.row(row['id']), 'approved')
        report = self.teacher.call('report', {'period': self.period})
        lecture = next(r for r in report['rows'] if r['type_id'] == 1)
        self.assertEqual(lecture['fact'], '2.00')
        self.assertEqual(lecture['clock_hours'], '1.50')
        self.assertEqual(report['approved_minutes'], 90)

    def test_atomic_import_and_dry_run(self):
        rows = [self.record('09:00', '10:00'), self.record('10:00', '11:00')]
        result = self.teacher.call('import', {'rows': rows, 'dry_run': True})
        self.assertEqual(result['count'], 2)
        self.assertEqual(self.teacher.call('entries'), [])
        bad = rows + [self.record('10:30', '11:30')]
        with self.assertRaises(ApiError):
            self.teacher.call('import', {'rows': bad, 'dry_run': False})
        self.assertEqual(self.teacher.call('entries'), [])
        self.teacher.call('import', {'rows': rows, 'dry_run': False})
        self.assertEqual(len(self.teacher.call('entries')), 2)

    def test_period_close_and_open(self):
        row = self.create()
        request = {'period': self.period, 'closed': 1, 'reason': 'Контрольное закрытие'}
        with self.assertRaises(ApiError):
            self.admin.call('period.set', request)
        self.send(self.teacher, row, 'submitted')
        self.send(self.head, self.row(row['id']), 'approved')
        self.admin.call('period.set', request)
        with self.assertRaises(ApiError):
            self.create('12:00', '13:00')
        self.admin.call('period.set', {**request, 'closed': 0})
        self.create('12:00', '13:00')

    def test_password_change_revokes_session(self):
        self.teacher.call('password', {'old': 'DemoTime2026!', 'new': 'NewValidPassword2026!'})
        with self.assertRaises(ApiError) as error:
            self.teacher.call('me')
        self.assertEqual(error.exception.status, 401)
        new = Api(self.url)
        new.login('teacher', 'NewValidPassword2026!')
        self.assertEqual(new.user['login'], 'teacher')

    def test_backup_restore_and_lock(self):
        self.create()
        result = self.admin.call('backup')
        backup = self.store.path.parent / 'backups' / result['filename']
        self.assertEqual(verify_backup(backup)['integrity'], 'ok')
        target = self.store.path.parent / 'restored.sqlite'
        restore(backup, target)
        with Store(target).connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM entries').fetchone()[0], 1)
        with InstanceLock(target):
            with self.assertRaises(DomainError):
                restore(backup, target)

    def test_feedback_privacy_and_answer(self):
        result = self.teacher.call('feedback.save', {'subject': 'Вопрос', 'message': 'Тестовое обращение'})
        self.assertEqual(self.other.call('feedback'), [])
        self.admin.call('feedback.save', {'id': result['id'], 'answer': 'Тестовый ответ'})
        self.assertEqual(self.teacher.call('feedback')[0]['answer'], 'Тестовый ответ')

    def test_exports(self):
        self.create()
        report = self.teacher.call('report', {'period': self.period})
        for extension, method in [('csv', exports.to_csv), ('xlsx', exports.to_xlsx), ('docx', exports.to_docx)]:
            path = Path(self.temp.name) / ('Отчёт.' + extension)
            method(report, path)
            self.assertGreater(path.stat().st_size, 100)
        from openpyxl import load_workbook
        book = load_workbook(Path(self.temp.name) / 'Отчёт.xlsx')
        self.assertEqual(book.active['D4'].value, 'План')
        self.assertEqual(book['Сведения']['B4'].value, report['approved_minutes'])
        for row in book.active.iter_rows(min_row=5, min_col=4, max_col=8):
            for cell in row:
                self.assertEqual(cell.data_type, 'n', cell.coordinate)
        self.assertEqual(sum(cell[0].value for cell in book.active.iter_rows(min_row=5, min_col=4, max_col=4)),
                         sum(float(row['plan']) for row in report['rows']))
        self.assertTrue(any(row[0].value < 0 for row in book.active.iter_rows(min_row=5, min_col=6, max_col=6)))
        self.assertEqual(book['Записи']['F2'].data_type, 'n')
        self.assertEqual(book['Записи']['F2'].value, 90)
        from docx import Document
        document = Document(Path(self.temp.name) / 'Отчёт.docx')
        self.assertEqual([c.text for c in document.tables[0].rows[0].cells], exports.HEADERS)
        self.assertGreater(document.sections[0].page_width, document.sections[0].page_height)
        self.assertIn('Согласовано, мин: 0', [p.text for p in document.paragraphs])

    def test_plan_version_and_unit_protection(self):
        row = self.teacher.call('plans', {'period': self.period})[0]
        self.head.call('plan.save', {**row, 'amount': '25'})
        with self.assertRaises(ApiError):
            self.head.call('plan.save', {**row, 'amount': '30'})
        with self.assertRaises(ApiError):
            self.admin.call('catalog.save', {'catalog': 'worktypes', 'record': {'id': 1, 'name': 'Лекция', 'unit': 'clock', 'active': 1}})

    def test_quality_and_search(self):
        self.create()
        found = self.teacher.call('search', {'period': self.period, 'query': 'контрольное'})
        self.assertEqual(len(found), 1)
        self.assertEqual(self.other.call('search', {'period': self.period, 'query': 'контрольное'}), [])
        self.assertEqual(len(self.teacher.call('calendar', {'period': self.period})), 1)
        self.assertIsInstance(self.teacher.call('quality', {'period': self.period}), list)

    def test_economics_reproducible(self):
        result = calculate()
        self.assertEqual(result['records_month'], 2400)
        self.assertEqual(result['total_saved_hours'], Decimal('95'))
        self.assertEqual(result['initial'], Decimal('226700'))
        self.assertEqual(result['net_month'], Decimal('49550'))
        self.assertEqual(len(sensitivity()), 4)


if __name__ == '__main__':
    unittest.main(verbosity=2)
