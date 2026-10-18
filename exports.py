"""Локальный экспорт согласованных данных. Документы не выдаются за официальный табель."""
import csv
from pathlib import Path
from domain import csv_safe


HEADERS = ['Преподаватель', 'Вид работы', 'Единица плана', 'План', 'Согласованный факт', 'Отклонение', 'Астрономические часы', 'На согласовании, мин']


def report_rows(report):
    for row in report['rows']:
        unit = 'акад. ч' if row['unit'] == 'academic' else 'ч'
        yield [row['teacher'], row['worktype'], unit, row['plan'], row['fact'], row['deviation'], row['clock_hours'], row['submitted_minutes']]


def to_csv(report, path):
    with open(path, 'w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(HEADERS)
        writer.writerows([[csv_safe(v) for v in row] for row in report_rows(report)])
    return Path(path)


def to_xlsx(report, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = 'План и факт'
    sheet.append([report['settings']['organization']])
    sheet.append(['Период', report['period']])
    sheet.append([report['settings']['report_note']])
    sheet.append(HEADERS)
    for row in report_rows(report):
        values = [csv_safe(value) for value in row[:3]]
        values.extend(float(value) for value in row[3:7])
        values.append(int(row[7]))
        sheet.append(values)
    for cell in sheet[4]:
        cell.font = Font(bold=True, color='000000')
        cell.fill = PatternFill('solid', fgColor='E7E7E7')
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = Border(bottom=Side(style='thin', color='D9D9D9'))
    for row in sheet.iter_rows(min_row=5, min_col=4, max_col=8):
        for cell in row:
            cell.number_format = '0' if cell.column == 8 else '0.00'
    for index, width in enumerate([32, 30, 17, 12, 18, 16, 20, 22], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'D5'
    sheet.auto_filter.ref = f'A4:H{sheet.max_row}'
    sheet.print_title_rows = '1:4'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    details = book.create_sheet('Записи')
    details.append(['Дата', 'Начало', 'Окончание', 'Преподаватель', 'Вид работы', 'Минуты', 'Состояние', 'Содержание'])
    for row in report['entries']:
        details.append([int(row[key]) if key == 'minutes' else csv_safe(row[key])
                        for key in ('workday', 'start', 'end', 'teacher', 'worktype', 'minutes', 'status', 'title')])
    details.freeze_panes = 'A2'
    for index in range(1, 9):
        details.column_dimensions[get_column_letter(index)].width = 23
    info = book.create_sheet('Сведения')
    info.append(['Показатель', 'Значение'])
    info.append(['Период', report['period']])
    info.append(['Дата формирования', report['generated_at']])
    info.append(['Согласовано, мин', report['approved_minutes']])
    info.append(['На проверке, мин', report['submitted_minutes']])
    info.append(['Первичных записей', len(report['entries'])])
    info.append(['Единица интервалов', 'Минута'])
    info.append(['Правило факта', 'Только согласованные записи'])
    info.append(['Суммирование плана', 'Отдельно по виду работы и единице'])
    info.column_dimensions['A'].width = 34
    info.column_dimensions['B'].width = 64
    book.save(path)
    return Path(path)


def to_docx(report, path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    normal = document.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for style in ('Title', 'Heading 1', 'Heading 2'):
        document.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    title = document.add_paragraph('Отчёт об учёте рабочего времени', style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(report['settings']['organization'])
    document.add_paragraph('Период: ' + report['period'])
    document.add_paragraph(report['settings']['report_note'])
    table = document.add_table(rows=1, cols=len(HEADERS))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [4.1, 4.2, 2.5, 2, 3.5, 2.6, 3.4, 3.4]
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    values = [HEADERS] + list(report_rows(report))
    for index, row_values in enumerate(values):
        row = table.rows[0] if index == 0 else table.add_row()
        row_property = row._tr.get_or_add_trPr()
        row_property.append(OxmlElement('w:cantSplit'))
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            row_property.append(repeat)
        for cell, value, width in zip(row.cells, row_values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.line_spacing = 1
            run = paragraph.add_run(str(value))
            run.font.size = Pt(10)
            run.bold = index == 0
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'EAEAEA')
                cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph('Сводные показатели', style='Heading 1')
    document.add_paragraph('Согласовано, мин: ' + str(report['approved_minutes']))
    document.add_paragraph('На проверке, мин: ' + str(report['submitted_minutes']))
    document.add_paragraph('Строк плана и факта: ' + str(len(report['rows'])))
    document.add_paragraph('Первичных записей в доступной области: ' + str(len(report['entries'])))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    document.add_paragraph('В фактический результат включены только согласованные записи. Академические и астрономические часы не суммируются между собой.')
    document.add_paragraph('Дата формирования: ' + report['generated_at'])
    document.save(path)
    return Path(path)


def read_import(path):
    with open(path, encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream, delimiter=';')
        expected = {'teacher_id', 'workday', 'start', 'end', 'type_id', 'title', 'evidence'}
        if set(reader.fieldnames or []) != expected:
            raise ValueError('Колонки CSV: teacher_id;workday;start;end;type_id;title;evidence')
        rows = list(reader)
    if not 0 < len(rows) <= 1000:
        raise ValueError('Файл должен содержать от 1 до 1000 строк')
    return rows


def import_template(path):
    with open(path, 'w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(['teacher_id', 'workday', 'start', 'end', 'type_id', 'title', 'evidence'])
    return Path(path)
